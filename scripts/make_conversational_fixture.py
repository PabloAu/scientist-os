"""Build original CC0 mixed-format sources; numbers are fictional, not studies."""
import argparse
import json
from pathlib import Path
import shutil


def build(root: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen.canvas import Canvas
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise ValueError("Select an empty folder; fixture generation never overwrites source files")
    def text(name, value):
        (root / name).write_text(value, encoding="utf-8", newline="\n")
    text("README.md", "# Fictional phantom experiment\n\nOriginal CC0-1.0 teaching material. No experiment occurred. All study reports and values are invented. Do not cite as real literature.\n")
    text("planning.md", "# Planned phantom experiment\n\nQuestion: does a correction reduce residual drift in a microscopy phantom? Plan: compare three independent preparations per condition, two technical measurements per preparation, 10 ms exposure. Materials: synthetic phantom beads, lot SIM-001; simulated buffer lot BUF-001. Protocol PH-1.0. No biological endpoint.\n\nCompeting explanations: correction removes drift; acquisition exposure differs; preparation imbalance; technical repeat noise. Controls proposed: randomize acquisition order, matched exposure, fixed fiducial reference, calibration and blinded QC. These are plans requiring a scientist's choice in real research.\n")
    text("protocol.md", "# Protocol PH 1.0\n\nFictional steps: prepare independent phantom batches; acquire technical measurements; measure residual displacement in nm; retain failures. Planned exposure 10 ms. Material lot SIM-001, storage room temperature (fictional). Actual execution and deviations are controlled by execution-notes.md, not this plan.\n")
    text("execution-notes.md", "# Actual execution description for the fixture\n\nThis describes the invented fixture only; no physical experiment occurred. The independent unit is preparation, not measurement. A1 A2 A3 belong to uncorrected; B1 B2 B3 belong to corrected. Conditions use distinct preparations and are not paired. Two measurements were planned per preparation; A2 and B2 each have one measurement because the second acquisition was missing, not excluded for its value. Exposure was 12 ms in both groups, differing from the 10 ms plan. All supplied numeric measurements are retained. Residual units are nm. Sample counts are three preparations and five observed technical measurements per group. No randomization or blinding was actually established.\n")
    shutil.copyfile(Path(__file__).resolve().parents[1] / "examples/pipelines/fictional_measurements.csv", root / "measurements.csv")
    text("study-sheets.md", "# Fictional study sheets\n\nAll are invented arithmetic exercises; no literature review is claimed.\n\nReport A sentence 1: fictional mean difference 0 arbitrary units, standard error 1, independent simulation A.\nReport B sentence 1: fictional mean difference 2 arbitrary units, standard error 1, independent simulation B.\nReport C sentence 1: fictional mean difference 4 arbitrary units, standard error 1, independent simulation C.\nReport B duplicate: same simulation B and numbers; retain as duplicate, do not pool twice.\nReport D: ratio 1.3 with standard error 0.2 on a ratio scale; exclude from a mean-difference synthesis.\n\nA-C share the generated estimand, direction, scale and units by construction. Protocol: include only independent mean-difference reports; DL random effects with fixed-effect, leave-one-out and omit-C sensitivity. No real populations, publication bias assessment or causal validity. Search scope: these five finite fixture reports only, 2026-09-07.\n")
    (root / "unsupported-instrument.bin").write_bytes(b"INERT FICTIONAL FORMAT; no instrument parser; retain a gap\n")
    # The original figure is a direct plotted representation of declared fixture means.
    canvas = Canvas(str(root / "fictional-paper.pdf"), pagesize=letter, invariant=1)
    canvas.setTitle("Fictional phantom methods note")
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(54, 730, "Fictional phantom methods note")
    canvas.setFont("Helvetica", 11)
    lines = ["Original CC0 teaching fixture. Not a publication or an experiment.",
             "Question: compare residual drift under two invented conditions.",
             "Independent unit: preparation. Three per condition, five measurements each.",
             "Equal-weight preparation means: uncorrected 8 nm; corrected 4 nm.",
             "Figure 1 shows descriptive means only. It does not establish efficacy."]
    for i, line in enumerate(lines):
        canvas.drawString(54, 692 - i * 20, line)
    canvas.setStrokeColor(colors.HexColor("#243d54"))
    canvas.line(125, 300, 480, 300)
    canvas.line(125, 300, 125, 530)
    for value in (0, 2, 4, 6, 8, 10):
        canvas.drawString(96, 297 + value * 22, str(value))
    canvas.setFillColor(colors.HexColor("#376787"))
    for x, value, label in ((205, 8, "Uncorrected"), (365, 4, "Corrected")):
        canvas.rect(x - 25, 300, 50, value * 22, fill=1, stroke=0)
        canvas.setFillColor(colors.black)
        canvas.drawCentredString(x, 275, label)
        canvas.drawCentredString(x, 315 + value * 22, str(value) + " nm")
        canvas.setFillColor(colors.HexColor("#376787"))
    canvas.setFillColor(colors.black)
    canvas.drawString(54, 234, "Figure 1. Known fictional preparation means; no uncertainty shown here.")
    canvas.drawString(54, 211, "Potential explanations include acquisition, batch variation and correction.")
    canvas.drawString(54, 76, "Source authority: generated fixture. Original figure must be inspected visually.")
    canvas.save()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.3333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Fictional phantom experiment"
    box = slide.shapes.add_textbox(Inches(.65), Inches(1.4), Inches(11.8), Inches(1))
    box.text_frame.text = "Three independent preparations per condition. Technical measurements are nested."
    box.text_frame.paragraphs[0].font.size = Pt(22)
    data = CategoryChartData()
    data.categories = ["Uncorrected", "Corrected"]
    data.add_series("Residual drift (nm)", (8, 4))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2.7), Inches(10.8), Inches(3.8), data)
    slide.notes_slide.notes_text_frame.text = "Original CC0 fictional fixture; editable chart of preparation means. No error bars on this source slide; the later analysis must estimate uncertainty. Not biological validation."
    deck.save(root / "lab-context.pptx")
    text("fixture-manifest.json", json.dumps({"license": "CC0-1.0", "synthetic": True,
         "source_owner": "Scientist OS fixture authoring agent", "scientific_approval": False,
         "actual_experiment": False}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    build(parser.parse_args().output)
