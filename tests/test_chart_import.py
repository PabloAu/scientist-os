import io
import zipfile

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from scientist_os.publishing import import_document
from scientist_os.workspace import Workspace


def chart_deck():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Known fixture chart"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("value", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                          Inches(1), Inches(2), Inches(5), Inches(3), data)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def test_native_chart_workbook_is_inert_but_does_not_block_slide_text(tmp_path):
    record = import_document(Workspace(tmp_path), filename="chart.pptx", data=chart_deck())
    assert "Known fixture chart" in record["content"]
    assert any("values are not extracted" in s for s in record["metadata"]["extraction"]["limitations"])


def test_embedded_chart_workbook_cannot_hide_a_macro(tmp_path):
    original = zipfile.ZipFile(io.BytesIO(chart_deck()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for name in original.namelist():
            payload = original.read(name)
            if name.endswith(".xlsx"):
                nested = io.BytesIO()
                with zipfile.ZipFile(nested, "w") as workbook:
                    workbook.writestr("[Content_Types].xml", "<Types/>")
                    workbook.writestr("xl/vbaProject.bin", b"inert test macro marker")
                payload = nested.getvalue()
            target.writestr(name, payload)
    with pytest.raises(ValueError, match="Macro"):
        import_document(Workspace(tmp_path), filename="chart.pptx", data=output.getvalue())
