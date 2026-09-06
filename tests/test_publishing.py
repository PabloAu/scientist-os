"""Scientific export round trips and hostile document boundaries."""

import hashlib
import io
import zipfile

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
import pytest

from scientist_os.publishing import (
    MAX_IMPORT_BYTES, attachment_bytes, create_presentation, export_manuscript,
    export_presentation, import_document, update_presentation, create_document_excerpt,
    publication_figure,
)
from scientist_os.service import analyze
from scientist_os.workspace import Workspace


@pytest.fixture
def workspace(tmp_path):
    return Workspace(tmp_path / "project")


@pytest.fixture
def evidence(workspace):
    dataset = workspace.create_record("dataset", "Registered teaching measurements", "group,day,value\nControl,c1,10\nControl,c2,12\nTreatment,t1,13\nTreatment,t2,15\n", metadata={"units": "nm", "synthetic": True})
    result = analyze(workspace, {"dataset_id": dataset["id"], "method": "summary", "value_column": "value", "group_column": "group", "unit_column": "day"})
    reference = workspace.create_record("reference", "Fictional teaching reference", metadata={"authors": ["A. Example"], "year": 2026, "journal": "Fictional methods", "doi": "10.0000/example"})
    return dataset, result["analysis"], result["output"], reference


def _archive(parts):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts:
            archive.writestr(name, value)
    return data.getvalue()


def _docx(text="A measured result."):
    document = Document()
    document.add_heading("Proposal", 0)
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx():
    presentation = Presentation()
    for text in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_import_preserves_original_and_reopens(workspace):
    original = b"Scientific proposal\nTwo units, not twenty independent samples.\n"
    record = import_document(workspace, filename="proposal.md", data=original, category="proposal")
    assert record["kind"] == "document"
    assert record["metadata"]["external_allowed"] is False
    assert record["metadata"]["attachment_sha256"] == hashlib.sha256(original).hexdigest()
    assert "[Lines 1-2]" in record["content"]
    assert record["metadata"]["extraction"]["segment_count"] == 1
    reopened = Workspace(workspace.root)
    assert attachment_bytes(reopened, record["id"]) == (original, "proposal.md", "text/markdown")
    assert not workspace.audit()


def test_original_bytes_are_not_silently_redacted(workspace):
    original = b"api example sk-aaaaaaaaaaaaaaaaaaaaaaaaa"
    record = import_document(workspace, filename="notes.txt", data=original)
    assert record["metadata"]["original_contains_unredacted_bytes"] is True
    assert attachment_bytes(workspace, record["id"])[0] == original


def test_attachment_change_detected(workspace):
    record = import_document(workspace, filename="notes.txt", data=b"abcd")
    path = workspace.root / "attachments" / record["metadata"]["attachment_sha256"]
    path.write_bytes(b"dcba")
    with pytest.raises(RuntimeError, match="digest"):
        attachment_bytes(workspace, record["id"])
    with pytest.raises(RuntimeError, match="digest"):
        import_document(workspace, filename="same.txt", data=b"abcd")


def test_attachment_missing_detected(workspace):
    record = import_document(workspace, filename="notes.txt", data=b"abcd")
    (workspace.root / "attachments" / record["metadata"]["attachment_sha256"]).unlink()
    with pytest.raises(RuntimeError, match="missing"):
        attachment_bytes(workspace, record["id"])


@pytest.mark.parametrize("filename,data", [
    ("../paper.txt", b"content"), ("C:\\paper.txt", b"content"), ("a.docm", b"PK"),
    ("a.docx", b"not a zip"), ("a.pdf", b"not a PDF"), ("a.txt", b"\xff"),
    ("a.txt", b"bad\x00text"), ("a.txt", b""),
], ids=["traversal", "windows-path", "macro", "bad-zip", "bad-pdf", "bad-utf8", "null", "empty"])
def test_invalid_imports_do_not_register(workspace, filename, data):
    with pytest.raises(ValueError):
        import_document(workspace, filename=filename, data=data)
    assert workspace.list_records() == []
    assert not (workspace.root / "attachments").exists()


def test_oversized_import_rejected(workspace):
    with pytest.raises(ValueError, match="20 MB"):
        import_document(workspace, filename="large.txt", data=b"x" * (MAX_IMPORT_BYTES + 1))
    assert workspace.list_records() == []


@pytest.mark.parametrize("bad_name,bad_xml", [
    ("../escaped", b"x"), ("/absolute", b"x"), ("x\\escape", b"x"),
    ("word/vbaProject.bin", b"x"), ("word/embeddings/oleObject1.bin", b"x"),
    ("word/activeX/control.xml", b"<x/>"),
    ("word/document.xml", b'<!DOCTYPE root [<!ENTITY x SYSTEM "file:///etc/passwd">]><root>&x;</root>'),
    ("word/document.xml", b'<!DOCTYPE root [<!ENTITY a "x"><!ENTITY b "&a;&a;">]><root>&b;</root>'),
], ids=["traversal", "absolute", "backslash", "macro", "object", "activex", "xxe", "entity"])
def test_unsafe_office_members_rejected(workspace, bad_name, bad_xml):
    data = _archive([("[Content_Types].xml", "<Types/>"), (bad_name, bad_xml)])
    with pytest.raises(ValueError):
        import_document(workspace, filename="document.docx", data=data)
    assert workspace.list_records() == []


def test_compression_bomb_rejected(workspace):
    data = _archive([("[Content_Types].xml", "<Types/>"), ("word/document.xml", "x" * 100000)])
    with pytest.raises(ValueError, match="compression"):
        import_document(workspace, filename="document.docx", data=data)


def test_docx_import_locators_and_tables(workspace):
    document = Document()
    document.add_paragraph("Paragraph one")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table cell evidence"
    output = io.BytesIO()
    document.save(output)
    record = import_document(workspace, filename="paper.docx", data=output.getvalue(), category="paper")
    assert "[Paragraph 1]\nParagraph one" in record["content"]
    assert "[Paragraph 2]\nTable cell evidence" in record["content"]
    assert "not page numbers" in " ".join(record["metadata"]["extraction"]["limitations"])


def test_pptx_import_respects_actual_slide_order(workspace):
    original = _pptx()
    parts = []
    with zipfile.ZipFile(io.BytesIO(original)) as source:
        for name in source.namelist():
            data = source.read(name)
            if name == "ppt/presentation.xml":
                from lxml import etree
                root = etree.fromstring(data)
                ids = root.find("{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst")
                ids.insert(0, ids[-1])
                data = etree.tostring(root)
            parts.append((name, data))
    record = import_document(workspace, filename="slides.pptx", data=_archive(parts))
    assert "[Slide 1]\nSecond slide" in record["content"]
    assert "[Slide 2]\nFirst slide" in record["content"]


def test_blank_pdf_is_retained_without_fabricated_text(workspace):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    record = import_document(workspace, filename="scan.pdf", data=output.getvalue())
    assert "[Page 1]" in record["content"]
    assert any("No readable text" in x for x in record["metadata"]["extraction"]["limitations"])


def test_encrypted_pdf_rejected(workspace):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="Encrypted"):
        import_document(workspace, filename="scan.pdf", data=output.getvalue())


def _slides_for(figure, reference):
    return [{"title": "Teaching experiment progress", "layout": "title", "body": "Synthetic examples for software verification"},
            {"title": "Independent-unit means preserve the sampling structure", "body": "Control and treatment measurements are averaged within independent units. These fictional values do not establish a biological effect.",
             "notes": "Discuss the next experimental design with the team.", "figure_id": figure["id"], "reference_ids": [reference["id"]]}]


def test_editable_pptx_preserves_values_notes_units_and_lineage(workspace, evidence):
    dataset, analysis, figure, reference = evidence
    record = create_presentation(workspace, title="Weekly research discussion", slides=_slides_for(figure, reference), source_ids=[dataset["id"]])
    output = export_presentation(workspace, record["id"])
    deck = Presentation(io.BytesIO(output))
    assert len(deck.slides) == 2
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, abs=1e-5)
    assert "unreviewed" in deck.core_properties.subject
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "nm" in text and "11" in text and "14" in text
    assert "one standard error, not a confidence interval" in text
    notes = deck.slides[1].notes_slide.notes_text_frame.text
    assert "Fictional teaching reference" in notes and analysis["id"] in notes
    assert "11.0, 10.0, 12.0" in notes
    assert "14.0, 13.0, 15.0" in notes
    assert dataset["id"] in notes
    with zipfile.ZipFile(io.BytesIO(output)) as archive:
        assert not any("media/image" in name for name in archive.namelist())  # plots are editable shapes
    assert workspace.get_record(record["id"])["review_status"] == "unreviewed"


def test_stale_input_blocks_pptx_export_and_update_is_revision_locked(workspace, evidence):
    dataset, _, figure, reference = evidence
    record = create_presentation(workspace, title="Research", slides=_slides_for(figure, reference))
    updated = update_presentation(workspace, record["id"], expected_revision=1, title="Research update", slides=record["metadata"]["slides"])
    assert updated["metadata"]["slides"][0]["id"] == record["metadata"]["slides"][0]["id"]
    with pytest.raises(RuntimeError, match="Stale"):
        update_presentation(workspace, record["id"], expected_revision=1, title="Old title", slides=record["metadata"]["slides"])
    workspace.update_record(dataset["id"], expected_revision=1, content=dataset["content"] + "Control,c3,9\n")
    with pytest.raises(RuntimeError):
        export_presentation(workspace, record["id"])


@pytest.mark.parametrize("slide", [
    {"title": "Dense", "body": "word " * 190}, {"title": "x" * 111},
    {"title": "Invalid", "layout": "unsupported"}, {"title": "Invalid", "notes": "x\x00y"},
    {"title": "Overflow", "body": "\n" * 15}, {"title": "Invalid", "unknown": True},
])
def test_bad_or_overflowing_slides_fail_before_save(workspace, slide):
    with pytest.raises(ValueError):
        create_presentation(workspace, title="Deck", slides=[slide])
    assert workspace.list_records() == []


def test_arbitrary_svg_cannot_be_embedded(workspace):
    figure = workspace.create_record("output", "Imported SVG", '<svg onload="alert(1)"></svg>', metadata={"format": "svg"})
    with pytest.raises(ValueError, match="reproducible"):
        create_presentation(workspace, title="Deck", slides=[{"title": "Figure", "figure_id": figure["id"]}])


def _manuscript(workspace, evidence):
    _, _, figure, reference = evidence
    sections = [{"id": "1" * 32, "title": "Results", "text": "The registered control and treatment unit means were 11 and 14 nm. <script>alert('x')</script>",
                 "figure_ids": [figure["id"]], "reference_ids": [reference["id"]], "supplementary": False},
                {"id": "2" * 32, "title": "Acquisition settings", "text": "Synthetic example. No experimental acquisition was performed.",
                 "figure_ids": [], "reference_ids": [], "supplementary": True}]
    return workspace.create_record("manuscript", "Experimental report", metadata={"sections": sections, "input_revisions": {figure["id"]: figure["revision"], reference["id"]: reference["revision"]}}, links=[figure["id"], reference["id"]])


def test_docx_contains_full_manuscript_figures_refs_supplements(workspace, evidence):
    record = _manuscript(workspace, evidence)
    output = export_manuscript(workspace, record["id"])
    doc = Document(io.BytesIO(output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "11 and 14 nm" in text
    assert "Supplementary · Acquisition settings" in text
    assert "[1] A. Example" in text and "Fictional teaching reference" in text
    assert "Human review: unreviewed" in text
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0]._inline.docPr.get("descr")
    assert "one standard error" in text


def test_html_escape_and_self_contained_figure(workspace, evidence):
    record = _manuscript(workspace, evidence)
    output = export_manuscript(workspace, record["id"], format="html").decode()
    assert "<script>" not in output
    assert "&lt;script&gt;" in output
    assert "data:image/png;base64," in output
    assert "Content-Security-Policy" in output
    assert "Fictional teaching reference" in output


@pytest.mark.parametrize("format", ["docx", "html"])
def test_stale_reference_blocks_manuscript_export(workspace, evidence, format):
    record = _manuscript(workspace, evidence)
    reference = evidence[-1]
    workspace.update_record(reference["id"], expected_revision=reference["revision"], title="Corrected paper title")
    with pytest.raises(RuntimeError):
        export_manuscript(workspace, record["id"], format=format)


def test_missing_embedded_revision_binding_rejected(workspace, evidence):
    record = _manuscript(workspace, evidence)
    metadata = dict(record["metadata"])
    metadata["input_revisions"] = {}
    workspace.update_record(record["id"], expected_revision=1, metadata=metadata)
    with pytest.raises(RuntimeError, match="binding"):
        export_manuscript(workspace, record["id"])


def test_excerpt_is_exact_private_revision_bound_and_reopenable(workspace):
    document = import_document(workspace, filename="paper.txt", data="Observed 🔬 signal.\nUncertain interpretation.".encode())
    text = document["content"]
    selected = "🔬 signal."
    start = len(text[:text.index(selected)].encode("utf-16-le")) // 2
    end = start + len(selected.encode("utf-16-le")) // 2
    excerpt = create_document_excerpt(workspace, document["id"], expected_revision=1, start=start, end=end, selected_text=selected, title="Selected observation")
    assert excerpt["kind"] == "source" and excerpt["content"] == selected
    assert excerpt["metadata"]["external_allowed"] is False
    assert excerpt["metadata"]["locators"] == ["Lines 1-2"]
    assert excerpt["metadata"]["input_revisions"] == {document["id"]: 1}
    assert Workspace(workspace.root).validate_current(excerpt["id"])["content"] == selected
    workspace.update_record(document["id"], expected_revision=1, content="Changed source")
    with pytest.raises(RuntimeError):
        workspace.validate_current(excerpt["id"])


@pytest.mark.parametrize("change", ["mismatch", "half-surrogate", "out-of-range", "stale", "boolean", "permission"])
def test_excerpt_rejects_invalid_selection_without_new_record(workspace, change):
    document = import_document(workspace, filename="paper.txt", data="🔬 observed".encode())
    text = document["content"]
    start = text.index("🔬")
    spec = {"expected_revision": 1, "start": start, "end": start + 2, "selected_text": "🔬", "title": "Excerpt"}
    if change == "mismatch":
        spec["selected_text"] = "fabricated"
    elif change == "half-surrogate":
        spec["start"] += 1
    elif change == "out-of-range":
        spec["end"] = 10000
    elif change == "stale":
        spec["expected_revision"] = 2
    elif change == "boolean":
        spec["start"] = True
    elif change == "permission":
        spec["external_allowed"] = "yes"
    with pytest.raises((ValueError, RuntimeError)):
        create_document_excerpt(workspace, document["id"], **spec)
    assert len(workspace.list_records()) == 1


def test_excerpt_checks_original_digest_and_explicit_disclosure(workspace):
    document = import_document(workspace, filename="paper.txt", data=b"Observed signal")
    start = document["content"].index("Observed")
    spec = {"expected_revision": 1, "start": start, "end": start + len("Observed signal"), "selected_text": "Observed signal", "title": "Observation", "external_allowed": True}
    excerpt = create_document_excerpt(workspace, document["id"], **spec)
    assert excerpt["metadata"]["external_allowed"] is True
    path = workspace.root / "attachments" / document["metadata"]["attachment_sha256"]
    path.write_bytes(b"Changed signal ")
    with pytest.raises(RuntimeError):
        create_document_excerpt(workspace, document["id"], **spec)


def test_large_excerpt_keeps_exact_text_without_exceeding_quote_limit(workspace):
    text = "measured " * 2000
    document = import_document(workspace, filename="paper.txt", data=text.encode())
    start = document["content"].index("measured")
    excerpt = create_document_excerpt(workspace, document["id"], expected_revision=1, start=start, end=start + len(text), selected_text=text, title="Extended methods")
    assert excerpt["content"] == text
    assert excerpt["metadata"]["citations"] == []
    assert excerpt["metadata"]["exact_excerpt_sha256"] == hashlib.sha256(text.encode()).hexdigest()


def test_pdf_javascript_and_external_decoder_filters_are_rejected(workspace):
    from pypdf.generic import DecodedStreamObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject()
    stream.set_data(b"dummy")
    stream[NameObject("/Filter")] = NameObject("/JBIG2Decode")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="unsupported stream filter"):
        import_document(workspace, filename="malicious.pdf", data=output.getvalue())
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.add_js("app.alert('untrusted')")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="active"):
        import_document(workspace, filename="active.pdf", data=output.getvalue())


@pytest.mark.parametrize("encoding", ["flate", "runlength"])
def test_pdf_expansion_is_bounded_before_text_extraction(workspace, encoding, monkeypatch):
    import zlib
    from pypdf import PageObject, filters
    from pypdf.generic import EncodedStreamObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    stream = EncodedStreamObject()
    if encoding == "flate":
        stream._data = zlib.compress(b"A" * (4 * 1024 * 1024 + 1))
        stream[NameObject("/Filter")] = NameObject("/FlateDecode")
    else:
        stream._data = b"\x81A" * (4 * 1024 * 1024 // 128 + 1) + b"\x80"
        stream[NameObject("/Filter")] = NameObject("/RunLengthDecode")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    calls = []
    monkeypatch.setattr(PageObject, "extract_text", lambda *args, **kwargs: calls.append(True) or "")
    original_limit = filters.ZLIB_MAX_OUTPUT_LENGTH
    with pytest.raises(ValueError, match="decoding limit"):
        import_document(workspace, filename="compressed.pdf", data=output.getvalue())
    assert not calls
    assert filters.ZLIB_MAX_OUTPUT_LENGTH == original_limit
    assert not workspace.list_records()


def test_docx_has_no_template_title_rule_or_blue_caption(workspace, evidence):
    record = _manuscript(workspace, evidence)
    output = export_manuscript(workspace, record["id"])
    doc = Document(io.BytesIO(output))
    assert not doc.styles.element.xpath(".//w:pBdr")
    assert str(doc.styles["Caption"].font.color.rgb) == "000000"


def test_publication_figure_is_png_from_current_registered_evidence(workspace, evidence):
    from PIL import Image
    dataset, _, figure, _ = evidence
    png = publication_figure(workspace, figure["id"])
    assert png.startswith(b"\x89PNG\r\n\x1a\n")
    with Image.open(io.BytesIO(png)) as image:
        assert image.format == "PNG"
        assert image.width == 1200
        assert 450 <= image.height <= 1000
    workspace.update_record(dataset["id"], expected_revision=1, content=dataset["content"] + "Control,c3,9\n")
    with pytest.raises(RuntimeError):
        publication_figure(workspace, figure["id"])
