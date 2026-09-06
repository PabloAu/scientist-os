"""Local document reading and evidence-linked Office exports.

Imported originals are inert attachments, not executable templates. Text is
untrusted source material; extraction never means that a paper was read or that
its claims were verified. All generated plots consume registered numeric results.
"""

from __future__ import annotations

import base64
from contextlib import contextmanager
import hashlib
import html
import io
import math
import os
from functools import lru_cache
from pathlib import Path, PurePosixPath
import re
import textwrap
import threading
from uuid import uuid4
import zipfile

from defusedxml import ElementTree

from .science import render_figure
from .workspace import Workspace

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_TEXT_BYTES = 1_500_000
_HASH = re.compile(r"[0-9a-f]{64}\Z")
_SID = re.compile(r"[0-9a-f]{32}\Z")
_PDF_LOCK = threading.RLock()
_MEDIA = {
    ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv",
}


def _text(value, field, limit, *, empty=True):
    if not isinstance(value, str) or len(value) > limit or (not empty and not value.strip()):
        raise ValueError(f"{field} must be text of at most {limit} characters")
    try:
        value.encode("utf-8")
    except UnicodeError as exc:
        raise ValueError(f"{field} must contain valid Unicode") from exc
    if any((ord(c) < 32 and c not in "\n\r\t") or ord(c) in {0xfffe, 0xffff} for c in value):
        raise ValueError(f"{field} contains unsupported control characters")
    return value


def _ids(value, field, limit=64):
    if not isinstance(value, list) or len(value) > limit or any(not isinstance(x, str) for x in value):
        raise ValueError(f"{field} must be an array of at most {limit} record IDs")
    if len(set(value)) != len(value):
        raise ValueError(f"{field} must not contain duplicate IDs")
    return value


def _safe_archive(data):
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        members = archive.infolist()
        if len(members) > 3000 or sum(m.file_size for m in members) > MAX_ARCHIVE_BYTES:
            raise ValueError("Office archive exceeds member or expanded-size limits")
        names = set()
        for member in members:
            name = member.filename
            path = PurePosixPath(name)
            if (name in names or "\\" in name or ":" in name or path.is_absolute()
                    or ".." in path.parts or member.flag_bits & 1):
                raise ValueError("Office archive contains unsafe, duplicate or encrypted members")
            names.add(name)
            if (member.file_size > 10 * 1024 * 1024
                    or member.file_size > max(member.compress_size, 1) * 250):
                raise ValueError("Office archive exceeds per-member or compression-ratio limits")
            if ((member.external_attr >> 16) & 0o170000) == 0o120000:
                raise ValueError("Office archive contains symbolic links")
            lower = name.lower()
            if any(x in lower for x in ("vbaproject", "/activex/", "/embeddings/")):
                raise ValueError("Macro, ActiveX and embedded-object Office documents are unsupported")
        if "[Content_Types].xml" not in names:
            raise ValueError("Not an Office Open XML document")
        # Parse each XML part through a DTD/entity-safe parser, including parts we
        # do not extract. Nothing from the archive is ever written to disk.
        xml = {}
        xml_bytes = 0
        node_count = 0
        for name in names:
            if name.lower().endswith((".xml", ".rels")):
                payload = archive.read(name)
                xml_bytes += len(payload)
                if xml_bytes > 12 * 1024 * 1024:
                    raise ValueError("Office XML exceeds the 12 MB parsing limit")
                xml[name] = ElementTree.fromstring(payload)
                node_count += sum(1 for _ in xml[name].iter())
                if node_count > 200_000:
                    raise ValueError("Office XML exceeds the element-count limit")
        content_types = archive.read("[Content_Types].xml").lower()
        if b"macroenabled" in content_types or b"vbaproject" in content_types:
            raise ValueError("Macro-enabled Office documents are unsupported")
        return xml
    except (zipfile.BadZipFile, OSError, KeyError, RuntimeError) as exc:
        raise ValueError("Invalid or unsupported Office archive") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unsafe or malformed Office XML") from exc


def _check_pdf_objects(reader):
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject
    seen = set()
    count = 0
    decoded_bytes = 0
    filters = {"/FlateDecode", "/Fl", "/ASCIIHexDecode", "/AHx", "/ASCII85Decode", "/A85",
               "/LZWDecode", "/LZW", "/RunLengthDecode", "/RL", "/DCTDecode", "/DCT", "/JPXDecode", "/CCITTFaxDecode", "/CCF"}

    def visit(value, depth=0):
        nonlocal count, decoded_bytes
        count += 1
        if count > 50_000 or depth > 50:
            raise ValueError("PDF object structure exceeds the supported extraction limits")
        if isinstance(value, IndirectObject):
            identity = (value.idnum, value.generation)
            if identity in seen:
                return
            seen.add(identity)
            visit(value.get_object(), depth + 1)
        elif isinstance(value, DictionaryObject):
            if any(key in value for key in ("/JavaScript", "/JS", "/EmbeddedFiles", "/RichMediaContent")):
                raise ValueError("PDF contains active or embedded content; provide a flattened PDF or text export")
            if isinstance(value, StreamObject):
                declared = value.get("/Filter", [])
                declared = declared if isinstance(declared, (ArrayObject, list)) else [declared]
                if any(str(item) not in filters for item in declared):
                    raise ValueError("PDF uses an unsupported stream filter; provide a text export")
                # Images are retained without pixel decoding. Decode text/font
                # streams once under the filter limits and cap their combined
                # cache before extracting any text.
                if value.get("/Subtype") != "/Image":
                    decoded_bytes += len(value.get_data())
                    if decoded_bytes > 32 * 1024 * 1024:
                        raise ValueError("PDF decoded text/font streams exceed the 32 MB total limit")
            for child in value.values():
                visit(child, depth + 1)
        elif isinstance(value, (ArrayObject, list)):
            for child in value:
                visit(child, depth + 1)

    visit(reader.trailer)


@contextmanager
def _pdf_limits():
    """Bound pypdf decoding before allocation, including initial object streams.

These are upstream's documented/configurable module limits. Serializing this
small extraction path prevents concurrent requests from changing the limits.
They are restored for other code using pypdf in the same interpreter.
"""
    from pypdf import filters
    limits = {name: 4 * 1024 * 1024 for name in (
        "MAX_DECLARED_STREAM_LENGTH", "MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH",
        "LZW_MAX_OUTPUT_LENGTH", "RUN_LENGTH_MAX_OUTPUT_LENGTH", "ZLIB_MAX_OUTPUT_LENGTH",
        "FLATE_MAX_ROW_LENGTH", "FLATE_MAX_BUFFER_SIZE",
    )}
    limits["ZLIB_MAX_RECOVERY_INPUT_LENGTH"] = 1024 * 1024
    with _PDF_LOCK:
        old = {key: getattr(filters, key) for key in limits}
        try:
            for key, value in limits.items():
                setattr(filters, key, min(old[key], value))
            yield
        finally:
            for key, value in old.items():
                setattr(filters, key, value)


def _extract(filename, data):
    suffix = PurePosixPath(filename).suffix.lower()
    segments = []
    limitations = ["Extracted text is untrusted source material, not verified evidence."]
    if suffix in {".docx", ".pptx"}:
        xml = _safe_archive(data)
        if suffix == ".docx":
            root = xml.get("word/document.xml")
            if root is None:
                raise ValueError("The file has no Word document body")
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for i, paragraph in enumerate(root.iter(namespace + "p"), 1):
                value = "".join(element.text or "" for element in paragraph.iter(namespace + "t"))
                if value.strip():
                    segments.append((f"Paragraph {i}", value))
            limitations += ["Paragraph numbers are extraction locators, not page numbers.",
                            "Formatting, figures, equations, comments, tracked deletions, headers and footnotes are not extracted."]
        else:
            names = sorted((name for name in xml if re.fullmatch(r"ppt/slides/slide[0-9]+\.xml", name)),
                           key=lambda name: int(re.search(r"([0-9]+)\.xml", name).group(1)))
            if not names or len(names) > 250:
                raise ValueError("Presentation must contain 1 to 250 slides")
            # Slide filenames need not reflect presentation order; resolve actual
            # order through the presentation relationship IDs.
            presentation = xml.get("ppt/presentation.xml")
            relationships = xml.get("ppt/_rels/presentation.xml.rels")
            if presentation is None or relationships is None:
                raise ValueError("Presentation order metadata is missing")
            rels = {r.attrib.get("Id"): r.attrib.get("Target", "") for r in relationships}
            ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            ordered = []
            for item in presentation.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"):
                target = rels.get(item.attrib.get(ns), "")
                if target.startswith("/"):
                    path = target.lstrip("/")
                else:
                    path = "ppt/" + target
                if path not in names:
                    raise ValueError("Unsupported presentation slide relationship")
                ordered.append(path)
            if len(ordered) != len(names) or len(set(ordered)) != len(ordered):
                raise ValueError("Presentation slide order is inconsistent")
            for i, name in enumerate(ordered, 1):
                paragraphs = ["".join(e.text or "" for e in paragraph.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
                              for paragraph in xml[name].iter("{http://schemas.openxmlformats.org/drawingml/2006/main}p")]
                segments.append((f"Slide {i}", "\n".join(x for x in paragraphs if x)))
            limitations += ["Slide text only: images, charts, equations, animation and speaker notes are not extracted."]
        limitations += ["External links and Office objects are never fetched or executed."]
        method = "Office XML text extraction"
    elif suffix == ".pdf":
        from pypdf import PdfReader
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF header is missing")
        try:
            with _pdf_limits():
                reader = PdfReader(io.BytesIO(data), strict=True)
                if reader.is_encrypted:
                    raise ValueError("Encrypted PDFs are unsupported; provide a readable copy")
                _check_pdf_objects(reader)
                if len(reader.pages) > 250:
                    raise ValueError("PDF exceeds 250 pages")
                total = 0
                for i, page in enumerate(reader.pages, 1):
                    # The decoder limits above apply before this check, including
                    # the object streams needed to parse the document catalog.
                    contents = page.get_contents()
                    if contents is not None and len(contents.get_data()) > 4 * 1024 * 1024:
                        raise ValueError("PDF page content exceeds the extraction limit")
                    value = page.extract_text() or ""
                    total += len(value.encode("utf-8"))
                    if total > MAX_TEXT_BYTES:
                        raise ValueError("Extracted document text exceeds 1.5 MB")
                    segments.append((f"Page {i}", value))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF text extraction failed or exceeded a decoding limit; provide a text export") from exc
        limitations += ["Native text layers only; scanned pages and images need a separate OCR workflow.",
                        "Reading order, symbols, tables and equations may be incomplete or rearranged."]
        method = "pypdf native text extraction"
    else:
        try:
            value = data.decode("utf-8-sig")
        except UnicodeError as exc:
            raise ValueError("Text imports must use UTF-8 encoding") from exc
        # Chunk line locators keep tables and long prose readable without storing
        # a second unbounded copy in metadata.
        lines = value.splitlines()
        segments = [(f"Lines {i + 1}-{min(i + 80, len(lines))}", "\n".join(lines[i:i + 80]))
                    for i in range(0, len(lines), 80)]
        method = "UTF-8 text extraction"
    if len(segments) > 10_000:
        raise ValueError("Document exceeds 10,000 extraction segments")
    text = "\n\n".join(f"[{locator}]\n{value}" for locator, value in segments)
    _text(text, "Extracted text", MAX_TEXT_BYTES)
    if len(text.encode("utf-8")) > MAX_TEXT_BYTES:
        raise ValueError("Extracted document text exceeds 1.5 MB")
    if not any(value.strip() for _, value in segments):
        limitations.append("No readable text was found. The original is retained; do not infer its contents.")
    return text, {"method": method, "limitations": limitations, "segment_count": len(segments)}


def _attachment_path(workspace, digest):
    if not isinstance(digest, str) or not _HASH.fullmatch(digest):
        raise ValueError("Invalid attachment digest")
    directory = workspace.root / "attachments"
    if directory.is_symlink() or directory.resolve() != workspace.root / "attachments":
        raise ValueError("Attachment directory must stay inside the workspace")
    directory.mkdir(exist_ok=True)
    path = directory / digest
    if path.is_symlink() or path.resolve() != directory.resolve() / digest:
        raise ValueError("Attachment path must stay inside the workspace")
    return path


def import_document(workspace: Workspace, *, filename: str, data: bytes, title: str = "", category: str = "document") -> dict:
    """Preserve original bytes and register extracted, explicitly private text."""
    _text(filename, "Filename", 180, empty=False)
    if any(x in filename for x in ("/", "\\", ":")) or filename in {".", ".."}:
        raise ValueError("Provide a filename, not a filesystem path")
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in _MEDIA:
        raise ValueError("Supported imports: PDF, DOCX, PPTX, TXT, MD and CSV")
    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_IMPORT_BYTES:
        raise ValueError("Document bytes must be nonempty and at most 20 MB")
    _text(title, "Title", 250)
    if not isinstance(category, str) or category not in {"document", "paper", "proposal", "presentation", "protocol", "notes"}:
        raise ValueError("Unsupported library category")
    text, extraction = _extract(filename, data)
    digest = hashlib.sha256(data).hexdigest()
    path = _attachment_path(workspace, digest)
    if path.exists():
        if path.stat().st_size != len(data) or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise RuntimeError("Existing attachment bytes do not match their digest")
    else:
        temporary = path.with_name(uuid4().hex + ".tmp")
        try:
            with temporary.open("xb") as stream:
                stream.write(data)
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
    return workspace.create_record("document", title.strip() or filename, text, metadata={
        "category": category, "filename": filename, "media_type": _MEDIA[suffix],
        "attachment_sha256": digest, "attachment_bytes": len(data), "extraction": extraction,
        "external_allowed": False, "original_contains_unredacted_bytes": True,
    })


def attachment_bytes(workspace: Workspace, record_id: str) -> tuple[bytes, str, str]:
    record = workspace.get_record(record_id)
    if record["kind"] != "document":
        raise ValueError("Select an imported document")
    metadata = record["metadata"]
    path = _attachment_path(workspace, metadata.get("attachment_sha256"))
    if not path.exists():
        raise RuntimeError("Original attachment is missing; restore the complete workspace backup")
    if path.stat().st_size != metadata.get("attachment_bytes") or path.stat().st_size > MAX_IMPORT_BYTES:
        raise RuntimeError("Original attachment size has changed")
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != metadata["attachment_sha256"]:
        raise RuntimeError("Original attachment digest has changed")
    filename = re.sub(r'[^\w. ()-]', "_", str(metadata.get("filename", "document")))[:180]
    media_type = _MEDIA.get(PurePosixPath(filename).suffix.lower(), "application/octet-stream")
    return data, filename, media_type


def create_document_excerpt(workspace: Workspace, record_id: str, *, expected_revision: int,
                            start: int, end: int, selected_text: str, title: str,
                            external_allowed: bool = False) -> dict:
    """Register a bounded exact extract for explicit model context selection."""
    _text(title, "Excerpt title", 250, empty=False)
    _text(selected_text, "Selected text", 20_000, empty=False)
    if type(expected_revision) is not int or expected_revision < 1:
        raise ValueError("expected_revision must be a positive integer")
    if type(start) is not int or type(end) is not int or not 0 <= start < end:
        raise ValueError("Selection offsets must be ordered nonnegative UTF-16 integers")
    if type(external_allowed) is not bool:
        raise ValueError("External model permission must be an explicit boolean")
    with workspace.transaction():
        record = workspace.validate_current(record_id)
        if record["kind"] != "document":
            raise ValueError("Select an imported project document")
        if record["revision"] != expected_revision:
            raise RuntimeError("Document changed; select the excerpt again from its current revision")
        attachment_bytes(workspace, record_id)
        encoded = record["content"].encode("utf-16-le")
        if end * 2 > len(encoded):
            raise ValueError("Selection extends beyond the document")
        try:
            prefix = encoded[:start * 2].decode("utf-16-le")
            selection = encoded[start * 2:end * 2].decode("utf-16-le")
        except UnicodeError as exc:
            raise ValueError("Selection splits a UTF-16 character") from exc
        if selection != selected_text:
            raise ValueError("Selected text no longer matches the document exactly")
        markers = list(re.finditer(r"(?m)^\[(Page \d+|Slide \d+|Paragraph \d+|Lines \d+-\d+)\]$", record["content"]))
        first = max((i for i, marker in enumerate(markers) if marker.start() <= len(prefix)), default=0)
        locators = [marker.group(1) for marker in markers[first:] if marker.start() < len(prefix) + len(selection)]
        return workspace.create_record("source", title, selection, metadata={
            "source_type": "document_excerpt", "document_id": record_id,
            "selection_utf16": {"start": start, "end": end}, "locators": locators[:500],
            "original_attachment_sha256": record["metadata"]["attachment_sha256"],
            "extraction": record["metadata"].get("extraction", {}),
            "input_revisions": {record_id: record["revision"]},
            "citations": [{"record_id": record_id, "quote": selection, "sha256": record["sha256"]}]
            if len(selection.encode("utf-8")) <= 16_384 else [],
            "exact_excerpt_sha256": hashlib.sha256(selection.encode("utf-8")).hexdigest(),
            "external_allowed": external_allowed,
        }, links=[record_id])


def _slides(slides):
    if not isinstance(slides, list) or not 1 <= len(slides) <= 40:
        raise ValueError("A presentation needs 1 to 40 slides")
    result = []
    for item in slides:
        if not isinstance(item, dict) or set(item) - {"id", "title", "body", "notes", "figure_id", "reference_ids", "layout"}:
            raise ValueError("Unsupported slide fields")
        sid = item.get("id", uuid4().hex)
        if not isinstance(sid, str) or not _SID.fullmatch(sid):
            raise ValueError("Slide ID must be a 32-character lowercase UUID hex string")
        layout = item.get("layout", "evidence")
        if not isinstance(layout, str) or layout not in {"title", "section", "evidence", "two_column"}:
            raise ValueError("Unsupported slide layout")
        title = _text(item.get("title"), "Slide title", 110, empty=False)
        body = _text(item.get("body", ""), "Slide body", 300 if layout in {"title", "section"} else 900)
        notes = _text(item.get("notes", ""), "Speaker notes", 8000)
        figure_id = item.get("figure_id") or None
        if figure_id is not None and not isinstance(figure_id, str):
            raise ValueError("figure_id must be a record ID or null")
        if figure_id and layout in {"title", "section"}:
            raise ValueError("Use an evidence or two-column layout for figures")
        references = _ids(item.get("reference_ids", []), "Slide references", 12)
        # Layout planning rejects dense text before it is persisted. Never shorten
        # a scientific claim to make it fit without the author's intervention.
        if layout in {"title", "section"}:
            _fit_text(title, 11.5, 2.2, 38)
            _fit_text(body, 10.8, 1.8, 22)
        else:
            _fit_text(title, 11.8, 1.1, 28)
            if figure_id:
                _fit_text(body, 4.5, 4.5, 19)
            elif layout == "two_column":
                for column in _columns(body):
                    _fit_text(column, 5.6, 4.6, 22)
            else:
                _fit_text(body, 11.75, 4.6, 24)
        result.append({"id": sid, "title": title, "body": body, "notes": notes,
                       "figure_id": figure_id, "reference_ids": references, "layout": layout})
    if len({s["id"] for s in result}) != len(result):
        raise ValueError("Slide IDs must be unique")
    return result


def _wrapped(value, width, maximum, label):
    lines = []
    for paragraph in value.split("\n"):
        lines += textwrap.wrap(paragraph, width=width, break_long_words=True, break_on_hyphens=False) or [""]
    if len(lines) > maximum:
        raise ValueError(f"{label} is too dense for this layout; split it across slides or shorten it")
    return "\n".join(lines)


@lru_cache(maxsize=24)
def _font(size):
    from matplotlib import get_data_path
    from PIL import ImageFont
    # A bundled metric font makes the fit check portable. Conservative padding
    # accommodates Office's Aptos substitution; no OS font discovery or downloads.
    return ImageFont.truetype(str(Path(get_data_path()) / "fonts/ttf/DejaVuSans.ttf"), size * 4)


def _fit_text(value, width, height, size):
    font = _font(size)
    maximum = width * 72 * 4 / 1.08
    lines = []
    for paragraph in value.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word) if line else word
            if font.getlength(candidate) <= maximum:
                line = candidate
                continue
            if line:
                lines.append(line)
                line = ""
            while font.getlength(word) > maximum:
                end = 1
                while end < len(word) and font.getlength(word[:end + 1]) <= maximum:
                    end += 1
                lines.append(word[:end])
                word = word[end:]
            line = word
        lines.append(line)
    if len(lines) * size * 1.22 > height * 72:
        raise ValueError("Slide text is too dense for its available area; split the slide or shorten the text")
    return lines


def _columns(body):
    blocks = body.split("\n\n")
    if len(blocks) >= 2:
        cut = min(range(1, len(blocks)), key=lambda i: abs(len("\n\n".join(blocks[:i])) - len("\n\n".join(blocks[i:]))))
        return "\n\n".join(blocks[:cut]), "\n\n".join(blocks[cut:])
    lines = body.split("\n")
    if len(lines) == 1:
        words = body.split()
        cut = math.ceil(len(words) / 2)
        return " ".join(words[:cut]), " ".join(words[cut:])
    cut = math.ceil(len(lines) / 2)
    return "\n".join(lines[:cut]), "\n".join(lines[cut:])


def _presentation_record(workspace, title, slides, source_ids):
    title = _text(title, "Presentation title", 250, empty=False)
    slides = _slides(slides)
    source_ids = _ids([] if source_ids is None else source_ids, "Presentation sources")
    linked = list(source_ids)
    for slide in slides:
        if slide["figure_id"]:
            _figure(workspace, slide["figure_id"], max_entries=8)
            linked.append(slide["figure_id"])
        for reference_id in slide["reference_ids"]:
            if workspace.get_record(reference_id)["kind"] != "reference":
                raise ValueError("Slide references must be reference records")
            linked.append(reference_id)
    linked = list(dict.fromkeys(linked))
    records = [workspace.validate_current(record_id) for record_id in linked]
    content = "\n\n".join(f"## {s['title']}\n{s['body']}\n\nSpeaker notes:\n{s['notes']}" for s in slides)
    metadata = {"studio_type": "presentation", "studio_version": 1, "slides": slides,
                "source_ids": source_ids, "external_allowed": False,
                "input_revisions": {r["id"]: r["revision"] for r in records}}
    return title, content, metadata, linked


def create_presentation(workspace: Workspace, *, title: str, slides: list[dict], source_ids: list[str] | None = None) -> dict:
    with workspace.transaction():
        title, content, metadata, links = _presentation_record(workspace, title, slides, source_ids)
        return workspace.create_record("presentation", title, content, metadata=metadata, links=links)


def update_presentation(workspace: Workspace, record_id: str, *, expected_revision: int, title: str,
                        slides: list[dict], source_ids: list[str] | None = None) -> dict:
    with workspace.transaction():
        old = workspace.get_record(record_id)
        if old["kind"] != "presentation":
            raise ValueError("Select a presentation record")
        if source_ids is None:
            source_ids = old["metadata"].get("source_ids", [])
        title, content, metadata, links = _presentation_record(workspace, title, slides, source_ids)
        metadata["external_allowed"] = old["metadata"].get("external_allowed", False)
        return workspace.update_record(record_id, expected_revision=expected_revision, title=title,
                                       content=content, metadata=metadata, links=links)


def _reference(record):
    if record["kind"] != "reference":
        raise ValueError("Bibliography entries must be reference records")
    metadata = record["metadata"]
    authors = metadata.get("authors", [])
    if not isinstance(authors, list) or len(authors) > 100 or any(not isinstance(x, str) for x in authors):
        raise ValueError("Reference authors must be an array of author names")
    for author in authors:
        _text(author, "Reference author", 300)
    for key in ("journal", "doi", "url"):
        _text(metadata.get(key, ""), f"Reference {key}", 4000)
    if metadata.get("year") is not None and type(metadata["year"]) is not int:
        raise ValueError("Reference year must be an integer or null")
    parts = [", ".join(authors), str(metadata.get("year") or "n.d."), record["title"],
             metadata.get("journal", ""), ("doi:" + metadata["doi"]) if metadata.get("doi") else metadata.get("url", "")]
    return ". ".join(str(x).strip().rstrip(".") for x in parts if x)


def _figure(workspace, record_id, *, max_entries=12):
    figure = workspace.validate_current(record_id)
    if figure["kind"] != "output" or figure["metadata"].get("format") != "svg":
        raise ValueError("Only registered reproducible scientific figures can be exported")
    analyses = [workspace.get_record(r) for r in figure["links"]]
    analysis = next((r for r in analyses if r["kind"] == "analysis" and "result" in r["metadata"]), None)
    if analysis is None or figure["metadata"].get("input_revisions", {}).get(analysis["id"]) != analysis["revision"]:
        raise ValueError("Figure has no revision-bound reproducible analysis")
    workspace.validate_current(analysis["id"])
    result = analysis["metadata"]["result"]
    title = figure["metadata"].get("figure_title", analysis["title"])
    svg = render_figure(result, title=title)  # validates all numeric geometry
    if svg != figure["content"]:
        raise RuntimeError("Registered figure differs from its reproducible result; regenerate it")
    entries = []
    if result["type"] == "csv_summary":
        for row in result["groups"]:
            if row["mean"] is None:
                raise ValueError("Figures with entirely missing groups need an explicit revised figure plan")
            error = row.get("standard_error")
            entries.append((str(row["group"]), row["mean"], row["mean"] - error if error is not None else row["mean"],
                            row["mean"] + error if error is not None else row["mean"]))
        measure = str(result.get("measurement_units", "Unresolved units"))
        caption = f"{result['method']}; {result['uncertainty']}. Units: {measure}."
    else:
        for row in result["studies"]:
            entries.append((str(row["study_id"]), row["effect"], *row["ci95"]))
        entries.append(("Pooled estimate", result["estimate"], *result["ci95"]))
        measure = str(result.get("effect_measure", "Unresolved effect measure"))
        caption = f"{result['method']}; 95% normal intervals, not prediction intervals. Effect: {measure}."
    if len(entries) > max_entries or any(len(e[0]) > 42 for e in entries) or len(measure) > 80:
        raise ValueError(f"This layout supports at most {max_entries} plot rows, labels up to 42 characters and units up to 80 characters; prepare a smaller figure")
    return {"record": figure, "analysis": analysis, "result": result, "title": title,
            "entries": entries, "measure": measure, "caption": caption, "svg": svg}


def _plot_png(figure):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    # FigureCanvas avoids pyplot global state in parallel local HTTP requests.
    plot = Figure(figsize=(8.0, 3.0 + len(figure["entries"]) * .27), dpi=150)
    FigureCanvasAgg(plot)
    axis = plot.add_subplot(111)
    for i, (label, value, low, high) in enumerate(figure["entries"]):
        axis.plot([low, high], [i, i], color="#245575", linewidth=1.5)
        axis.plot(value, i, "o", color="#245575", markersize=5)
    axis.set_yticks(range(len(figure["entries"])), [e[0] for e in figure["entries"]])
    axis.invert_yaxis()
    axis.set_xlabel(figure["measure"])
    axis.spines[["right", "top", "left"]].set_visible(False)
    axis.grid(axis="x", color="#dee5eb", linewidth=.6)
    axis.set_axisbelow(True)
    plot.tight_layout(pad=1.8)
    output = io.BytesIO()
    plot.savefig(output, format="png")
    return output.getvalue()


def publication_figure(workspace: Workspace, record_id: str) -> bytes:
    """Render current registered numeric evidence for authoring canvas previews."""
    with workspace.transaction():
        return _plot_png(_figure(workspace, record_id))


def _linked_record(workspace, owner, record_id):
    item = workspace.validate_current(record_id)
    if owner["metadata"].get("input_revisions", {}).get(record_id) != item["revision"]:
        raise RuntimeError("An embedded figure or reference has no current revision binding; save the authoring document again")
    return item


def _manuscript_snapshot(workspace, record_id):
    record = workspace.validate_current(record_id)
    if record["kind"] != "manuscript" or not isinstance(record["metadata"].get("sections"), list):
        raise ValueError("Select a section-based manuscript from the authoring studio")
    _text(record["title"], "Manuscript title", 1000, empty=False)
    sections = record["metadata"]["sections"]
    if not 1 <= len(sections) <= 40:
        raise ValueError("Manuscript must have 1 to 40 sections")
    refs, figures = {}, {}
    total = 0
    for section in sections:
        if not isinstance(section, dict):
            raise ValueError("Invalid manuscript section")
        _text(section.get("title"), "Section title", 300, empty=False)
        total += len(_text(section.get("text", ""), "Section text", 120_000))
        if total > 120_000 or type(section.get("supplementary", False)) is not bool:
            raise ValueError("Invalid manuscript section size or supplementary flag")
        for figure_id in _ids(section.get("figure_ids", []), "Section figures"):
            _linked_record(workspace, record, figure_id)
            figures[figure_id] = _figure(workspace, figure_id)
        for reference_id in _ids(section.get("reference_ids", []), "Section references"):
            reference = _linked_record(workspace, record, reference_id)
            refs[reference_id] = _reference(reference)
    return record, sections, refs, figures


def export_manuscript(workspace: Workspace, record_id: str, *, format: str = "docx") -> bytes:
    if not isinstance(format, str) or format not in {"docx", "html"}:
        raise ValueError("Manuscript export format must be docx or html")
    with workspace.transaction():
        record, sections, refs, figures = _manuscript_snapshot(workspace, record_id)
        numbers = {key: i for i, key in enumerate(refs, 1)}
        status = f"Revision {record['revision']} · Human review: {record['review_status']}"
        figure_numbers = {key: i for i, key in enumerate(figures, 1)}
        if format == "html":
            e = html.escape
            parts = ["<!doctype html><html lang='en'><meta charset='utf-8'>",
                     "<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; img-src data:; style-src 'unsafe-inline'\">",
                     f"<title>{e(record['title'])}</title><style>body{{font:17px/1.65 Georgia,serif;max-width:820px;margin:65px auto;padding:0 35px;color:#17212c}}h1,h2{{font-family:Arial,sans-serif;line-height:1.25}}p{{white-space:pre-wrap}}small,figcaption{{font:13px/1.5 Arial,sans-serif}}img{{max-width:100%}}figure{{margin:28px 0}}section{{margin-top:32px}}@media print{{body{{margin:20mm}}section{{break-inside:auto}}figure{{break-inside:avoid}}}}</style>",
                     f"<h1>{e(record['title'])}</h1><small>{e(status)}</small>"]
            for section in sections:
                prefix = "Supplementary · " if section.get("supplementary") else ""
                parts.append(f"<section><h2>{e(prefix + section['title'])}</h2><p>{e(section.get('text', ''))}</p>")
                if section.get("reference_ids"):
                    parts.append("<small>Section references: " + ", ".join(f"[{numbers[r]}]" for r in section["reference_ids"]) + "</small>")
                for figure_id in section.get("figure_ids", []):
                    figure = figures[figure_id]
                    encoded = base64.b64encode(_plot_png(figure)).decode("ascii")
                    parts.append(f"<figure><img alt='{e(figure['title'], quote=True)}' src='data:image/png;base64,{encoded}'><figcaption>Figure {figure_numbers[figure_id]}. {e(figure['title'])}. {e(figure['caption'])}</figcaption></figure>")
                parts.append("</section>")
            parts += ["<h2>References</h2><ol>"] + [f"<li>{e(value)}</li>" for value in refs.values()] + ["</ol></html>"]
            return "\n".join(parts).encode("utf-8")
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        document = Document()
        properties = document.core_properties
        properties.author = ""
        properties.last_modified_by = ""
        # OOXML core properties impose a 255-character limit. The full author
        # title remains intact in the document heading, never silently shortened.
        properties.title = record["title"] if len(record["title"]) <= 255 else ""
        properties.subject = status
        styles = document.styles
        styles["Normal"].font.name = "Cambria"
        styles["Normal"].font.size = Pt(11)
        styles["Normal"].paragraph_format.line_spacing = 1.15
        styles["Normal"].paragraph_format.space_after = Pt(8)
        for style in styles:
            for border in list(style.element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
        for name in ("Title", "Heading 1", "Heading 2"):
            styles[name].font.name = "Calibri"
            styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for name in ("Caption", "Subtitle"):
            styles[name].font.color.rgb = RGBColor(0, 0, 0)
            styles[name].font.bold = False
            styles[name].font.italic = False
        styles["Subtitle"].font.size = Pt(10)
        section_format = document.sections[0]
        section_format.top_margin = section_format.bottom_margin = Inches(.8)
        section_format.left_margin = section_format.right_margin = Inches(.85)
        document.add_heading(record["title"], 0)
        document.add_paragraph(status, style="Subtitle")
        for section in sections:
            document.add_heading(("Supplementary · " if section.get("supplementary") else "") + section["title"], 1)
            for paragraph in section.get("text", "").split("\n\n"):
                document.add_paragraph(paragraph)
            if section.get("reference_ids"):
                document.add_paragraph("Section references: " + ", ".join(f"[{numbers[r]}]" for r in section["reference_ids"]))
            for figure_id in section.get("figure_ids", []):
                figure = figures[figure_id]
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                shape = paragraph.add_run().add_picture(io.BytesIO(_plot_png(figure)), width=Inches(6.1))
                shape._inline.docPr.set("descr", figure["title"] + ". " + figure["caption"])
                document.add_paragraph(f"Figure {figure_numbers[figure_id]}. {figure['title']}. {figure['caption']}", style="Caption")
        document.add_heading("References", 1)
        for reference_id, text in refs.items():
            document.add_paragraph(f"[{numbers[reference_id]}] {text}")
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()


def export_presentation(workspace: Workspace, record_id: str) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    with workspace.transaction():
        record = workspace.validate_current(record_id)
        if record["kind"] != "presentation":
            raise ValueError("Select a presentation record")
        slides = _slides(record["metadata"].get("slides"))
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
        deck.core_properties.title = record["title"]
        deck.core_properties.author = ""
        deck.core_properties.last_modified_by = ""
        deck.core_properties.subject = f"Revision {record['revision']}; human review: {record['review_status']}"
        navy, ink, blue, gray = "10283F", "172A3B", "256387", "5A6B7A"

        def textbox(slide, value, x, y, width, height, size=20, color=ink, bold=False):
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            frame = box.text_frame
            frame.clear()
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = 0
            frame.margin_top = frame.margin_bottom = 0
            for index, line in enumerate(_fit_text(value, width, height, size)):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(color)
                paragraph.font.bold = bold
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = Pt(size * 1.22)
            return box

        def line(slide, x1, y1, x2, y2, color=blue, width=1.5):
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            shape.line.color.rgb = RGBColor.from_string(color)
            shape.line.width = Pt(width)
            effect = shape._element.find(".//" + qn("a:effectRef"))
            if effect is not None:
                effect.set("idx", "0")

        for index, spec in enumerate(slides, 1):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            dark = spec["layout"] in {"title", "section"}
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(navy if dark else "FFFFFF")
            title = spec["title"]
            if dark:
                textbox(slide, title, .8, 1.65, 11.5, 2.2, 38, "FFFFFF", True)
                textbox(slide, spec["body"], .84, 4.2, 10.8, 1.8, 22, "D9E4EE")
            else:
                textbox(slide, title, .72, .48, 11.8, 1.1, 28, ink, True)
                if spec["figure_id"]:
                    _linked_record(workspace, record, spec["figure_id"])
                    figure = _figure(workspace, spec["figure_id"], max_entries=8)
                    textbox(slide, spec["body"], .75, 1.8, 4.5, 4.5, 19)
                    entries = figure["entries"]
                    minimum, maximum = min(e[2] for e in entries), max(e[3] for e in entries)
                    span = maximum - minimum
                    if not math.isfinite(span):
                        raise ValueError("Plot scale exceeds supported range")
                    padding = span * .12 if span else max(abs(minimum) * .12, 1e-12)
                    minimum, maximum = minimum - padding, maximum + padding
                    def xpos(value):
                        return 8.0 + (value - minimum) / (maximum - minimum) * 3.25
                    for j in range(4):
                        value = minimum + (maximum - minimum) * j / 3
                        x = xpos(value)
                        line(slide, x, 2.0, x, 5.8, "E2E9EF", .6)
                        textbox(slide, f"{value:.4g}", x - .22, 5.92, .75, .32, 10, gray)
                    for row, (label, value, low, high) in enumerate(entries):
                        y = 2.3 + row * min(.49, 3.2 / max(len(entries) - 1, 1))
                        textbox(slide, label, 5.55, y - .12, 2.25, .5, 12)
                        line(slide, xpos(low), y, xpos(high), y)
                        for edge in (low, high):
                            line(slide, xpos(edge), y - .05, xpos(edge), y + .05)
                        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xpos(value) - .04), Inches(y - .04), Inches(.08), Inches(.08))
                        marker.fill.solid()
                        marker.fill.fore_color.rgb = RGBColor.from_string(blue)
                        marker.line.fill.background()
                        effect = marker._element.find(".//" + qn("a:effectRef"))
                        if effect is not None:
                            effect.set("idx", "0")
                        textbox(slide, f"{value:.5g}", 11.5, y - .1, 1.08, .3, 11, gray)
                    textbox(slide, figure["measure"], 7.4, 6.32, 5.1, .5, 12, gray)
                elif spec["layout"] == "two_column":
                    left, right = _columns(spec["body"])
                    textbox(slide, left, .75, 1.85, 5.6, 4.6, 22)
                    textbox(slide, right, 6.95, 1.85, 5.6, 4.6, 22)
                else:
                    textbox(slide, spec["body"], .75, 1.85, 11.75, 4.6, 24)
            references = [_linked_record(workspace, record, r) for r in spec["reference_ids"]]
            footer = "Sources: " + "; ".join(f"[{j}] {r['metadata'].get('year') or 'n.d.'}" for j, r in enumerate(references, 1)) if references else ""
            if spec["figure_id"]:
                figure = _figure(workspace, spec["figure_id"], max_entries=8)
                footer = (footer + " · " if footer else "") + figure["caption"]
            footer = _wrapped(footer, 146, 2, "Source footnote")
            textbox(slide, footer, .75, 6.89, 11.1, .46, 9, "D9E4EE" if dark else gray)
            textbox(slide, f"{index:02}", 12.0, 6.98, .5, .3, 10, "D9E4EE" if dark else gray)
            notes = [spec["notes"], f"Scientist OS presentation {record['id']} revision {record['revision']}. Human review: {record['review_status']}.",
                     "Review status is a recorded human decision, not independent scientific validation."]
            notes += [f"[{j}] {_reference(reference)}\nRecord {reference['id']}, revision {reference['revision']}."
                      for j, reference in enumerate(references, 1)]
            if spec["figure_id"]:
                notes += [f"Figure {figure['record']['id']} revision {figure['record']['revision']}; analysis {figure['analysis']['id']} revision {figure['analysis']['revision']}. {figure['caption']}",
                          "Numeric plot entries (label, estimate, lower bound, upper bound):",
                          *[repr(entry) for entry in figure["entries"]], *figure["result"].get("warnings", [])]
            for source_id in record["metadata"].get("source_ids", []):
                source = _linked_record(workspace, record, source_id)
                notes.append(f"Project source: {source['title']} ({source_id}, revision {source['revision']}).")
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
        output = io.BytesIO()
        deck.save(output)
        return output.getvalue()
